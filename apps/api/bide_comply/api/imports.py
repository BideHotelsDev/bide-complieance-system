"""Settings › Import data (design: Import from the register Sheet; CMS-06).

Upload the BH-CMP-002 gate workbook, the BH-CMP-001 register, or any CSV/XLSX. The wizard detects the template,
matches columns automatically where it can, asks about the rest, shows exactly what will happen, then copies the data
one way. Nothing in the workbook changes, and it can be re-run until go-live. Imported "To confirm" items stay open
until evidence is filed; "Done" with no evidence imports as "To confirm".
"""

import csv
import io
import re
import uuid
from datetime import date, datetime

from fastapi import APIRouter, Request
from pydantic import BaseModel
from sqlalchemy import select, update

from bide_comply import audit, clock
from bide_comply.auth import Current, actor_for
from bide_comply.db import ConnDep
from bide_comply.engine import actions as action_rules
from bide_comply.schema import action, evidence, gate, gate_item, import_job, person, pool, property_
from bide_comply.services.common import problem
from bide_comply.storage import storage

router = APIRouter(prefix="/api/imports", tags=["import"])

GATE_FIELDS = {
    "ref": ("Gate item ref", ["ref"]),
    "section": ("Section", ["section"]),
    "item": ("Item", ["item", "requirement"]),
    "class": ("Class", ["class"]),
    "legal_basis": ("Legal basis", ["legal basis", "standard"]),
    "evidence_required": ("Evidence required", ["evidence required"]),
    "applies": ("Applies (Yes / N/A)", ["applies"]),
    "status": ("Gate status", ["status"]),
    "owner": ("Owner → match to people", ["owner"]),
    "evidence_link": ("Evidence link", ["evidence filed", "evidence link", "file name"]),
    "notes": ("Notes", ["notes", "advisory"]),
}
ACTION_FIELDS = {
    "ref": ("Action ID", ["id", "ref"]),
    "type": ("Type", ["type"]),
    "gate_ref": ("Gate ref", ["gate ref", "gate"]),
    "source_ref": ("Source", ["source"]),
    "description": ("Description", ["description", "action"]),
    "priority": ("Priority", ["priority"]),
    "owner": ("Owner", ["owner"]),
    "raised_on": ("Raised", ["raised"]),
    "due_on": ("Due", ["due"]),
    "status": ("Status", ["status"]),
}
STATUS_MAP = {"done": "done", "in progress": "in_progress", "not started": "not_started", "to confirm": "to_confirm",
              "n/a": "na", "na": "na"}


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _read_workbook(data: bytes, name: str) -> dict[str, list[list]]:
    if name.lower().endswith(".csv"):
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        return {"Sheet1": rows}
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    return out


def _header_row(rows: list[list]) -> int:
    for i, r in enumerate(rows[:15]):
        cells = [_norm(c) for c in r]
        if "ref" in cells or "id" in cells or any(c.startswith("item") for c in cells):
            return i
    return 0


def _auto_map(headers: list[str], fields: dict) -> dict[str, str | None]:
    out: dict[str, str | None] = {}
    for h in headers:
        n = _norm(h)
        hit = None
        for key, (_, names) in fields.items():
            if any(n == x or n.startswith(x) for x in names) and key not in out.values():
                hit = key
                break
        out[h] = hit
    return out


def _sheet_kind(name: str, headers: list[str]) -> str | None:
    hs = {_norm(h) for h in headers}
    if "action" in _norm(name) or ({"priority", "due"} <= hs):
        return "actions"
    if "dashboard" in _norm(name):
        return None
    if {"ref", "section"} <= hs or {"ref", "item"} <= hs:
        return "gate"
    return None


def _property_for(conn, sheet: str) -> dict | None:
    for p in conn.execute(select(property_.c.id, property_.c.name, property_.c.code)).mappings():
        if _norm(p["name"]) in _norm(sheet) or _norm(sheet) in _norm(p["name"]):
            return dict(p)
    return None


class StartIn(BaseModel):
    evidence_id: uuid.UUID


@router.post("")
def start(body: StartIn, request: Request, pr: Current, conn: ConnDep) -> dict:
    pr.require("import.run")
    ev = conn.execute(select(evidence).where(evidence.c.id == body.evidence_id, evidence.c.uploaded_by == pr.person_id)).mappings().first()
    if not ev:
        raise problem(404, "not_found", "Upload the file first.")
    data = storage.get(ev["storage_key"])
    try:
        book = _read_workbook(data, ev["original_name"] or ev["file_name"])
    except Exception as e:  # noqa: BLE001 - a broken file is a user problem, not a server error
        raise problem(422, "unreadable", f"That file couldn't be read as a spreadsheet ({e.__class__.__name__}).") from None
    sheets = []
    for name, rows in book.items():
        if not rows:
            continue
        hi = _header_row(rows)
        headers = [str(h).strip() for h in rows[hi] if h not in (None, "")]
        kind = _sheet_kind(name, headers)
        data_rows = [r for r in rows[hi + 1:] if any(c not in (None, "") for c in r)]
        prop = _property_for(conn, name) if kind == "gate" else None
        sheets.append({"name": name, "kind": kind, "rows": len(data_rows), "header_row": hi, "headers": headers,
                       "property_id": str(prop["id"]) if prop else None, "property": prop["name"] if prop else None,
                       "example": {h: (str(v)[:80] if v is not None else None) for h, v in zip(headers, data_rows[0] if data_rows else [], strict=False)}})
    gate_sheets = [s for s in sheets if s["kind"] == "gate"]
    detected = "BH-CMP-002 gate" if gate_sheets else ("Actions log" if any(s["kind"] == "actions" for s in sheets) else "spreadsheet")
    mapping = {}
    if gate_sheets:
        mapping["gate"] = _auto_map(gate_sheets[0]["headers"], GATE_FIELDS)
    act = next((s for s in sheets if s["kind"] == "actions"), None)
    if act:
        mapping["actions"] = _auto_map(act["headers"], ACTION_FIELDS)
    jid = uuid.uuid4()
    conn.execute(import_job.insert().values(id=jid, kind="gate_workbook" if gate_sheets else "csv", file_name=ev["original_name"] or ev["file_name"],
                                            evidence_id=ev["id"], sheets=sheets, mapping=mapping, created_by=pr.person_id))
    audit.record(conn, actor_for(request, pr), category="record", action="import.start", summary=f"Started an import of {ev['original_name']}",
                 entity_type="import", entity_id=jid)
    return {"id": str(jid), "detected": detected}


def _job(conn, pr, jid: uuid.UUID) -> dict:
    pr.require("import.run")
    j = conn.execute(select(import_job).where(import_job.c.id == jid)).mappings().first()
    if not j:
        raise problem(404, "not_found", "Not found.")
    return dict(j)


def _rows_for(conn, j: dict, sheet: dict) -> list[dict]:
    data = storage.get(conn.execute(select(evidence.c.storage_key).where(evidence.c.id == j["evidence_id"])).scalar())
    book = _read_workbook(data, j["file_name"])
    rows = book[sheet["name"]]
    hi = sheet["header_row"]
    headers = [str(h).strip() if h is not None else "" for h in rows[hi]]
    out = []
    for r in rows[hi + 1:]:
        if not any(c not in (None, "") for c in r):
            continue
        out.append({headers[i]: r[i] for i in range(min(len(headers), len(r))) if headers[i]})
    return out


def _plan(conn, j: dict) -> dict:
    """What will happen, without changing anything (design: Checks before import)."""
    gmap = j["mapping"].get("gate", {})
    amap = j["mapping"].get("actions", {})
    inv_g = {v: k for k, v in gmap.items() if v}
    inv_a = {v: k for k, v in amap.items() if v}
    gate_count, to_confirm, done_no_ev, props = 0, 0, [], set()
    for s in j["sheets"]:
        if s["kind"] != "gate":
            continue
        props.add(s["property"] or s["name"])
        for r in _rows_for(conn, j, s):
            ref = r.get(inv_g.get("ref", ""))
            if not ref or not re.match(r"^[A-H]\d\d$", str(ref).strip()):
                continue
            gate_count += 1
            st = STATUS_MAP.get(_norm(r.get(inv_g.get("status", ""))), "not_started")
            ev = r.get(inv_g.get("evidence_link", ""))
            if st == "to_confirm":
                to_confirm += 1
            if st == "done" and not ev:
                done_no_ev.append(f"{s['property'] or s['name']} {ref}")
    actions, role_owners = 0, set()
    known = {_norm(n) for n in conn.execute(select(person.c.name)).scalars()} | {_norm(n.split()[0]) for n in conn.execute(select(person.c.name)).scalars()}
    first_ref = last_ref = None
    for s in j["sheets"]:
        if s["kind"] != "actions":
            continue
        for r in _rows_for(conn, j, s):
            ref = r.get(inv_a.get("ref", ""))
            if not ref:
                continue
            actions += 1
            first_ref = first_ref or ref
            last_ref = ref
            o = str(r.get(inv_a.get("owner", "")) or "").strip()
            if o and _norm(o) not in known:
                role_owners.add(o)
    unmatched = [h for h, v in gmap.items() if not v]
    return {"gate_items": gate_count, "properties": sorted(props), "actions": actions, "to_confirm": to_confirm,
            "done_without_evidence": done_no_ev, "role_owners": sorted(role_owners), "unmatched_columns": unmatched,
            "errors": 0, "action_range": f"{first_ref} to {last_ref}" if first_ref else None,
            "matched": sum(1 for v in gmap.values() if v), "columns": len(gmap)}


@router.get("/{jid}")
def get_job(jid: uuid.UUID, pr: Current, conn: ConnDep) -> dict:
    j = _job(conn, pr, jid)
    gate_sheet = next((s for s in j["sheets"] if s["kind"] == "gate"), None)
    fields = [{"key": k, "label": v[0]} for k, v in GATE_FIELDS.items()]
    return {"id": str(jid), "file_name": j["file_name"], "status": j["status"], "sheets": j["sheets"],
            "detected": "BH-CMP-002 gate" if gate_sheet else "spreadsheet",
            "gate_mapping": [{"column": h, "example": (gate_sheet["example"] if gate_sheet else {}).get(h), "field": v,
                              "matched": v is not None} for h, v in j["mapping"].get("gate", {}).items()],
            "fields": fields, "plan": _plan(conn, j), "result": j["result"], "owner_map": j["owner_map"],
            "status_mapping": [{"from": "Done", "to": "Done if evidence"}, {"from": "In progress", "to": "In progress"},
                               {"from": "Not started", "to": "Not started"}, {"from": "To confirm", "to": "To confirm"},
                               {"from": "N/A", "to": "N/A reason from Notes"}]}


class MapIn(BaseModel):
    gate: dict[str, str | None] | None = None
    actions: dict[str, str | None] | None = None
    owner_map: dict[str, str] | None = None  # role text -> person id


@router.put("/{jid}")
def set_mapping(jid: uuid.UUID, body: MapIn, pr: Current, conn: ConnDep) -> dict:
    j = _job(conn, pr, jid)
    mapping = dict(j["mapping"])
    if body.gate is not None:
        mapping["gate"] = {**mapping.get("gate", {}), **body.gate}
    if body.actions is not None:
        mapping["actions"] = {**mapping.get("actions", {}), **body.actions}
    vals: dict = {"mapping": mapping, "status": "review"}
    if body.owner_map is not None:
        vals["owner_map"] = {**j["owner_map"], **body.owner_map}
    conn.execute(update(import_job).where(import_job.c.id == jid).values(**vals))
    return {"ok": True}


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%d %b %Y", "%d %b", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            d = datetime.strptime(s, fmt).date()
            return d.replace(year=clock.today().year) if fmt == "%d %b" else d
        except ValueError:
            continue
    return None


@router.post("/{jid}/run")
def run(jid: uuid.UUID, request: Request, pr: Current, conn: ConnDep) -> dict:
    j = _job(conn, pr, jid)
    plan = _plan(conn, j)
    unmapped_roles = [o for o in plan["role_owners"] if o not in j["owner_map"]]
    gmap = {v: k for k, v in j["mapping"].get("gate", {}).items() if v}
    amap = {v: k for k, v in j["mapping"].get("actions", {}).items() if v}
    people = {_norm(n): i for i, n in conn.execute(select(person.c.id, person.c.name))}
    people |= {_norm(n.split()[0]): i for i, n in conn.execute(select(person.c.id, person.c.name))}
    created = updated_ = acts = 0
    for s in j["sheets"]:
        if s["kind"] != "gate" or not s["property_id"]:
            continue
        pid = uuid.UUID(s["property_id"])
        gid = conn.execute(select(gate.c.id).where(gate.c.property_id == pid)).scalar()
        if not gid:
            gid = uuid.uuid4()
            conn.execute(gate.insert().values(id=gid, property_id=pid))
        existing = {r.ref: r.id for r in conn.execute(select(gate_item.c.ref, gate_item.c.id).where(gate_item.c.gate_id == gid))}
        for i, r in enumerate(_rows_for(conn, j, s)):
            ref = str(r.get(gmap.get("ref", "")) or "").strip()
            if not re.match(r"^[A-H]\d\d$", ref):
                continue
            st = STATUS_MAP.get(_norm(r.get(gmap.get("status", ""))), "not_started")
            ev = r.get(gmap.get("evidence_link", ""))
            applies = _norm(r.get(gmap.get("applies", ""))) not in ("n/a", "na", "no")
            if st == "done" and not ev:
                st = "to_confirm"  # no evidence, not done
            vals = {"section": ref[0], "item": str(r.get(gmap.get("item", "")) or ref), "class": str(r.get(gmap.get("class", "")) or "B")[:1].upper(),
                    "legal_basis": r.get(gmap.get("legal_basis", "")), "evidence_required": r.get(gmap.get("evidence_required", "")),
                    "applies": applies, "status": "na" if not applies else st, "evidence_link": str(ev) if ev else None,
                    "notes": r.get(gmap.get("notes", "")), "na_reason": r.get(gmap.get("notes", "")) if not applies else None,
                    "owner_text": r.get(gmap.get("owner", "")), "sort": i + 1}
            owner = _norm(vals["owner_text"])
            if owner in people:
                vals["owner_id"] = people[owner]
            if ref in existing:
                conn.execute(update(gate_item).where(gate_item.c.id == existing[ref]).values(**vals))
                updated_ += 1
            else:
                conn.execute(gate_item.insert().values(id=uuid.uuid4(), gate_id=gid, ref=ref, **vals))
                created += 1
    have = set(conn.execute(select(action.c.ref)).scalars())
    codes = {r.code: r.id for r in conn.execute(select(property_.c.code, property_.c.id))}
    for s in j["sheets"]:
        if s["kind"] != "actions":
            continue
        for r in _rows_for(conn, j, s):
            ref = str(r.get(amap.get("ref", "")) or "").strip().upper()
            if not ref or ref in have or ref[:2] not in codes:
                continue
            prio = {"high": "high", "urgent": "high", "med": "med", "medium": "med", "soon": "med", "low": "low", "later": "low"}.get(
                _norm(r.get(amap.get("priority", ""))), "med")
            typ = {"advisory": "advisory", "work required": "work_required"}.get(_norm(r.get(amap.get("type", ""))), "action")
            raised = _date(r.get(amap.get("raised_on", ""))) or clock.today()
            owner_txt = str(r.get(amap.get("owner", "")) or "").strip()
            oid = people.get(_norm(owner_txt)) or (uuid.UUID(j["owner_map"][owner_txt]) if owner_txt in j["owner_map"] else None)
            conn.execute(action.insert().values(
                id=uuid.uuid4(), ref=ref, property_id=codes[ref[:2]], type=typ, source="manual",
                source_ref=r.get(amap.get("source_ref", "")), gate_ref=r.get(amap.get("gate_ref", "")),
                description=str(r.get(amap.get("description", "")) or ref), priority=prio, owner_id=oid,
                owner_text=None if oid else owner_txt or None, raised_on=raised,
                due_on=_date(r.get(amap.get("due_on", ""))) or action_rules.due_date(prio, raised),
                status="to_confirm" if _norm(r.get(amap.get("status", ""))) == "to confirm" else "open"))
            have.add(ref)
            acts += 1
    result = {"gate_created": created, "gate_updated": updated_, "actions": acts, "roles_left_to_map": unmapped_roles,
              "at": clock.now().isoformat()}
    conn.execute(update(import_job).where(import_job.c.id == jid).values(status="imported", result=result))
    audit.record(conn, actor_for(request, pr), category="record", action="import.run",
                 summary=f"Imported {created + updated_} gate items and {acts} Actions from {j['file_name']}",
                 entity_type="import", entity_id=jid, after=result)
    return result


@router.get("/options/owners")
def owner_options(pr: Current, conn: ConnDep) -> dict:
    pr.require("import.run")
    ppl = conn.execute(select(person.c.id, person.c.name).where(person.c.employment_status == "active").order_by(person.c.name)).all()
    pools = conn.execute(select(pool.c.id, pool.c.name).order_by(pool.c.name)).all()
    return {"people": [{"id": str(p.id), "name": p.name} for p in ppl], "pools": [{"id": str(p.id), "name": p.name} for p in pools]}
